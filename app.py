"""
Prisma - Lectura de entorno PESTEL participativa. Genérica: cada organización se configura con una plantilla
(ver plantillas.py); el código de las pantallas es el mismo para todas.

Ejecutar:
    pip install -r requirements.txt
    python app.py
Luego abrir http://127.0.0.1:5000  (formulario comunitario: http://127.0.0.1:5000/reportar)
"""
import base64
import csv
import hashlib
import io
import json
import os
import secrets
import uuid
from datetime import datetime
from functools import wraps

from flask import (Flask, Response, abort, flash, g, jsonify, redirect, render_template, request,
                   send_from_directory, session, url_for)
from werkzeug.security import check_password_hash

import analisis
import ia
from database import FACTORES, ahora, conectar, hora_local, inicializar, leer_ajuste, plantilla_activa

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 4 * 1024 * 1024  # las fotos llegan reducidas desde el celular
FOTOS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fotos")
os.makedirs(FOTOS_DIR, exist_ok=True)
inicializar()  # crea la base de datos con los datos de demostración si no existe


def clave_secreta():
    """Clave que firma las sesiones. Se toma de PRISMA_SECRET o se genera una al azar la primera vez y se guarda."""
    if os.environ.get("PRISMA_SECRET"):
        return os.environ["PRISMA_SECRET"]
    con = conectar()
    clave = leer_ajuste(con, "clave_sesion")
    if not clave:
        clave = secrets.token_hex(32)
        con.execute("INSERT OR REPLACE INTO ajustes (clave, valor) VALUES ('clave_sesion', ?)", (clave,))
        con.commit()
    con.close()
    return clave


app.secret_key = clave_secreta()
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")

# Nombre de la aplicación: se cambia solo aquí y se actualiza en todas las pantallas.
APP_NOMBRE = "Prisma"
APP_LEMA = "Lectura de entorno"

OPCIONES = {
    "tendencia": ["Sube", "Estable", "Baja"],
    "horizonte": ["Corto", "Mediano", "Largo"],
    "confiabilidad": ["Alta", "Media", "Baja"],
    "frecuencia": ["Diaria", "Semanal", "Mensual", "Trimestral"],
    "estado": ["Pendiente", "En análisis", "Aprobada", "Descartada"],
    "avance": ["Sin iniciar", "En curso", "Cumplido", "Atrasado"],
}
GRAVEDAD = {1: "Poco", 2: "Algo", 3: "Bastante", 4: "Mucho", 5: "Muy grave"}


# ---------------------------------------------------------------- plantilla de la organización
@app.before_request
def cargar_plantilla():
    g.P = plantilla_activa()


# ---------------------------------------------------------------- protección CSRF
# Cada formulario lleva un código secreto de la sesión; así otra página no puede enviar formularios en nombre del usuario.
# El formulario público de reporte y el ingreso no lo necesitan (no actúan con una sesión iniciada).
SIN_CSRF = {"login", "api_sugerir", "api_senal"}


@app.before_request
def proteger_formularios():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    if request.method == "POST" and request.endpoint not in SIN_CSRF:
        enviado = request.form.get("csrf") or request.headers.get("X-CSRF", "")
        if not secrets.compare_digest(enviado, session["csrf"]):
            abort(400)


@app.errorhandler(400)
def solicitud_invalida(_error):
    flash("La página llevaba mucho tiempo abierta. Vuelva a intentarlo.", "warning")
    return redirect(request.referrer or url_for("inicio"))


def T():
    return g.P["T"]


# ---------------------------------------------------------------- utilidades
def registrar(con, accion, detalle=""):
    con.execute("INSERT INTO bitacora (fecha, usuario, accion, detalle) VALUES (?, ?, ?, ?)",
                (ahora(), session.get("usuario", "formulario público"), accion, detalle))


def requiere_login(f):
    @wraps(f)
    def envoltura(*args, **kwargs):
        if "usuario" not in session:
            return redirect(url_for("login", siguiente=request.path))
        return f(*args, **kwargs)
    return envoltura


def requiere_rol(*roles):
    def decorador(f):
        @wraps(f)
        def envoltura(*args, **kwargs):
            if session.get("rol") not in roles:
                flash("Su rol no tiene permiso para esta acción.", "warning")
                return redirect(request.referrer or url_for("inicio"))
            return f(*args, **kwargs)
        return envoltura
    return decorador


def escenarios_activos():
    return session.get("escenarios", [])


def nombres_escenarios(con):
    ids = escenarios_activos()
    if not ids:
        return []
    return [r["nombre"] for r in con.execute(
        "SELECT nombre FROM escenarios WHERE id IN (%s)" % ",".join("?" * len(ids)), ids).fetchall()]


def contexto(municipio=None):
    con = conectar()
    return con, analisis.evaluar(con, escenarios_activos(), municipio), nombres_escenarios(con)


def puede_ver_sensible():
    return session.get("rol") in ("directivo", "analista")


def municipios(con):
    return [dict(m) for m in con.execute("SELECT * FROM municipios ORDER BY departamento, nombre").fetchall()]


@app.template_filter("minuscula")
def minuscula(texto):
    """Pone en minúscula solo la primera letra (respeta siglas como SENA o EUDR)."""
    if not texto:
        return ""
    pre = "¿" if texto.startswith("¿") else ""
    resto = texto[len(pre):]
    return pre + (resto[:1].lower() + resto[1:] if not resto[:2].isupper() else resto)


@app.template_filter("iniciales")
def iniciales(nombre):
    palabras = [p for p in (nombre or "").split() if p.lower() not in ("de", "del", "la", "el", "y")]
    return "".join(p[0] for p in palabras[:2]).upper()


@app.template_filter("dec")
def decimal_coma(valor):
    if valor is None:
        return "–"
    return f"{valor:.1f}".replace(".", ",")


@app.template_filter("fecha_corta")
def fecha_corta(texto):
    try:
        d = datetime.strptime(texto[:16], "%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return texto or ""
    meses = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
    return f"{d.day} {meses[d.month - 1]}"


@app.context_processor
def globales():
    activa = False
    if "usuario" in session:
        try:
            c = conectar()
            activa = ia.disponible(c)
            c.close()
        except Exception:
            pass
    P = getattr(g, "P", None) or plantilla_activa()
    opciones = {**OPCIONES, "cultivo": P["T"]["categorias"], "quien": P["T"]["quien"]}
    return {"IA_ACTIVA": activa, "APP": APP_NOMBRE, "LEMA": APP_LEMA, "ROLES": P["T"]["roles"], "FACTORES": FACTORES,
            "OPCIONES": opciones, "T": P["T"], "PLANTILLA": P, "ACTORES": P["actores"], "GRAVEDAD": GRAVEDAD, "hoy": hora_local().strftime("%d/%m/%Y")}


# ---------------------------------------------------------------- acceso
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        con = conectar()
        u = con.execute("SELECT * FROM usuarios WHERE usuario = ?", (request.form.get("usuario", "").strip(),)).fetchone()
        if u and check_password_hash(u["clave_hash"], request.form.get("clave", "")):
            escenarios = session.get("escenarios", [])
            session.clear()
            session.update(usuario=u["usuario"], nombre=u["nombre"], rol=u["rol"], escenarios=escenarios)
            registrar(con, "Inicio de sesión", T()["roles"][u["rol"]])
            con.commit()
            con.close()
            destino = request.args.get("siguiente", "")
            return redirect(destino if destino.startswith("/") and not destino.startswith("//") else url_for("inicio"))
        con.close()
        flash("Usuario o contraseña incorrectos.", "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------------------------------------------------------------- inicio y tablero
@app.route("/")
@requiere_login
def inicio():
    con, variables, nombres = contexto()
    factores = analisis.indices_factores(variables, "efectivo")
    for f, b, t, c in zip(factores, analisis.indices_factores(variables, "puntaje_base"),
                          analisis.indices_factores(variables, "puntaje"), analisis.indices_factores(variables, "com_puntaje")):
        f["base"] = b["indice"]
        f["delta"] = round((t["indice"] or 0) - (b["indice"] or 0), 1)
        f["tec"], f["tec_clase"] = t["indice"], t["clase"]
        f["com"], f["com_clase"] = c["indice"], c["clase"]
    mas, menos = analisis.brechas(variables)
    riesgos = sorted([v for v in variables if v["tipo"] == "R"], key=lambda v: -v["efectivo"])
    estados = {r["variable_id"]: r["estado"] for r in con.execute("SELECT * FROM decisiones").fetchall()}
    sugeridas = analisis.decisiones_sugeridas(variables)
    resumen = {
        "criticos": len([v for v in riesgos if v["nivel"] == "Crítico"]),
        "altos": len([v for v in riesgos if v["nivel"] == "Alto"]),
        "pendientes": len([v for v in sugeridas if estados.get(v["id"], "Pendiente") in ("Pendiente", "En análisis")]),
        "variables": len(variables),
        "escenarios": con.execute("SELECT COUNT(*) FROM escenarios").fetchone()[0],
        "por_validar": con.execute("SELECT COUNT(*) FROM senales WHERE estado = 'Por validar'").fetchone()[0],
        "senales": sum(v["com_n"] for v in variables),
        "municipios": con.execute("SELECT COUNT(*) FROM municipios").fetchone()[0],
        "compromisos": con.execute("SELECT COUNT(*) FROM decisiones WHERE estado = 'Aprobada'").fetchone()[0],
        "top": [v["nombre"] for v in riesgos[:2]],
    }
    por_id = {v["id"]: v for v in variables}
    decs = {r["variable_id"]: dict(r) for r in con.execute("SELECT * FROM decisiones").fetchall()}
    por_revisar = [v for v in sugeridas if decs.get(v["id"], {}).get("estado", "Pendiente") in ("Pendiente", "En análisis")]
    compromisos = [{**por_id[k], **d} for k, d in decs.items() if d["estado"] == "Aprobada" and k in por_id]
    emergentes = grupos_emergentes(con)
    entrando = analisis.escenarios_en_entrada(con)
    con.close()
    return render_template("inicio.html", entrando=entrando, factores=factores, r=resumen, brechas_mas=mas, brechas_menos=menos,
                           escenarios=nombres, urgentes=[v for v in riesgos if v["efectivo"] >= 10][:3],
                           por_revisar=por_revisar[:3], compromisos=compromisos, emergentes=emergentes)


@app.route("/tablero")
@requiere_login
def tablero():
    con, variables, nombres = contexto()
    factores = analisis.indices_factores(variables, "efectivo")
    tecnica = analisis.indices_factores(variables, "puntaje")
    comunidad = analisis.indices_factores(variables, "com_puntaje")
    riesgos = sorted([v for v in variables if v["tipo"] == "R"], key=lambda v: -v["efectivo"])
    oportunidades = sorted([v for v in variables if v["tipo"] == "O"], key=lambda v: -v["efectivo"])
    estados = {r["variable_id"]: r["estado"] for r in con.execute("SELECT * FROM decisiones").fetchall()}
    sugeridas = analisis.decisiones_sugeridas(variables)
    kpi = {
        "criticos": len([v for v in riesgos if v["nivel"] == "Crítico"]),
        "altos": len([v for v in riesgos if v["nivel"] == "Alto"]),
        "oportunidades": len(oportunidades),
        "pendientes": len([v for v in sugeridas if estados.get(v["id"], "Pendiente") == "Pendiente"]),
    }
    mas, menos = analisis.brechas(variables)
    con.close()
    return render_template("tablero.html", factores=factores, tecnica=tecnica, comunidad=comunidad, riesgos=riesgos,
                           oportunidades=oportunidades, kpi=kpi, escenarios=nombres,
                           alertas=[v for v in riesgos if v["efectivo"] >= 10], brechas_mas=mas, brechas_menos=menos)


@app.route("/revision", methods=["GET", "POST"])
@requiere_login
@requiere_rol("directivo", "analista")
def revision():
    """«¿Dónde podemos estar equivocados?»: la IA hace de abogado del diablo sobre la lectura PESTEL."""
    con = conectar()
    if request.method == "POST":
        variables = analisis.evaluar(con, ())
        resultado, motor = ia.revision_critica(con, variables)
        guardar_cache(con, "revision", resultado, motor)
        registrar(con, "Revisión crítica", "IA" if motor == "ia" else "modo básico")
        con.commit()
        if motor != "ia" and ia.disponible(con):
            flash("La IA no respondió; se muestra la revisión básica. " + (ia.ultimo_error["texto"] or ""), "warning")
        con.close()
        return redirect(url_for("revision"))
    ultima = leer_cache(con, "revision")
    con.close()
    return render_template("revision.html", r=ultima)


# ---------------------------------------------------------------- territorio
def resumen_municipio(con, m):
    variables = analisis.evaluar(con, escenarios_activos(), m["id"])
    con_senal = [v for v in variables if v["com_n"]]
    riesgos = sorted([v for v in con_senal if v["tipo"] == "R"], key=lambda v: -(v["com_puntaje"] or 0))
    maximo = riesgos[0]["com_puntaje"] if riesgos else None
    m = dict(m)
    m.update(senales=sum(v["com_n"] for v in variables), maximo=maximo,
             nivel=analisis.nivel(maximo), clase=analisis.CLASE_NIVEL[analisis.nivel(maximo)],
             principal=riesgos[0]["nombre"] if riesgos else None,
             por_validar=con.execute("SELECT COUNT(*) FROM senales WHERE municipio_id = ? AND estado = 'Por validar'",
                                     (m["id"],)).fetchone()[0])
    return m, variables


SQL_SENALES = """SELECT s.*, v.nombre AS variable, v.factor, COALESCE(v.sensible, 0) AS sensible,
                         COALESCE(v.icono, 'fa-lightbulb') AS icono, m.nombre AS municipio
                  FROM senales s LEFT JOIN variables v ON v.id = s.variable_id
                  JOIN municipios m ON m.id = s.municipio_id"""


def senal_dict(fila):
    d = dict(fila)
    try:
        d["ia"] = json.loads(d.get("ia") or "null")
    except ValueError:
        d["ia"] = None
    if not d.get("variable"):
        d["variable"] = d.get("emergente") or "Otro tema"
    return d


def leer_cache(con, clave):
    fila = con.execute("SELECT contenido, motor, fecha FROM ia_cache WHERE clave = ?", (clave,)).fetchone()
    if not fila:
        return None
    d = json.loads(fila["contenido"])
    if isinstance(d, dict):
        d.setdefault("_motor", fila["motor"])
        d.setdefault("_fecha", fila["fecha"])
    return d


def guardar_cache(con, clave, contenido, motor):
    con.execute("INSERT OR REPLACE INTO ia_cache (clave, contenido, motor, fecha) VALUES (?, ?, ?, ?)",
                (clave, json.dumps(contenido, ensure_ascii=False), motor, ahora()))


def emergentes_existentes(con):
    return [r[0] for r in con.execute("SELECT DISTINCT emergente FROM senales WHERE emergente IS NOT NULL "
                                      "AND variable_id IS NULL AND estado != 'Descartada'")]


def grupos_emergentes(con, ver_sensible=True):
    """Temas nuevos que la comunidad reporta y que todavía no son variables."""
    filas = con.execute(SQL_SENALES + " WHERE s.variable_id IS NULL AND s.emergente IS NOT NULL "
                        "AND s.estado != 'Descartada' ORDER BY s.fecha DESC").fetchall()
    grupos = {}
    for f in filas:
        g = grupos.setdefault(f["emergente"], {"tema": f["emergente"], "senales": [], "municipios": set(), "validadas": 0})
        g["senales"].append(senal_dict(f))
        g["municipios"].add(f["municipio"])
        g["validadas"] += f["estado"] == "Validada"
    salida = sorted(grupos.values(), key=lambda g: (-len(g["municipios"]), -len(g["senales"])))
    for g in salida:
        g["municipios"] = sorted(g["municipios"])
    return salida


def datos_boletin(con, variables, mid=None):
    """Resumen compacto (sin datos personales) que se le entrega a la IA para el boletín."""
    riesgos = sorted([v for v in variables if v["com_n"] and v["tipo"] == "R"], key=lambda v: -(v["com_puntaje"] or 0))
    sql = SQL_SENALES + " WHERE s.estado = 'Validada' AND COALESCE(v.sensible, 0) = 0 AND s.descripcion != ''"
    args = []
    if mid:
        sql += " AND s.municipio_id = ?"
        args.append(mid)
    voces = [f"{f['municipio']}: {f['descripcion']}" for f in con.execute(sql + " ORDER BY s.fecha DESC LIMIT 15", args)]
    compromisos = [f"{r['decision']} ({r['actores']}, {r['avance']})" for r in con.execute(
        "SELECT d.*, v.decision FROM decisiones d JOIN variables v ON v.id = d.variable_id WHERE d.estado = 'Aprobada'")]
    return {
        "preocupaciones": [{"tema": v["nombre"], "senales": v["com_n"], "gravedad": v["com_gravedad"],
                            "lectura_comunidad": v["com_puntaje"], "lectura_tecnica": v["puntaje"],
                            "accion": v["decision"], "actor": v["responsable"]} for v in riesgos[:5]],
        "brechas": [{"tema": v["nombre"], "brecha": v["brecha"]} for v in variables
                    if v["tipo"] == "R" and v["brecha"] is not None and abs(v["brecha"]) >= analisis.UMBRAL_BRECHA],
        "voces": voces,
        "temas_emergentes": [{"tema": g["tema"], "reportes": len(g["senales"])} for g in grupos_emergentes(con)
                             if not mid or any(s["municipio_id"] == mid for s in g["senales"])],
        "compromisos": compromisos,
    }


@app.route("/territorio")
@requiere_login
def territorio():
    con = conectar()
    lista = sorted([resumen_municipio(con, m)[0] for m in municipios(con)], key=lambda m: -(m["maximo"] or 0))
    nombres = nombres_escenarios(con)
    con.close()
    return render_template("territorio.html", lista=lista, escenarios=nombres)


@app.route("/municipio/<mid>")
@requiere_login
def municipio(mid):
    con = conectar()
    m = con.execute("SELECT * FROM municipios WHERE id = ?", (mid,)).fetchone()
    if not m:
        con.close()
        abort(404)
    m, variables = resumen_municipio(con, m)
    tecnica = analisis.indices_factores(variables, "puntaje")
    comunidad = analisis.indices_factores(variables, "com_puntaje")
    senales = [senal_dict(f) for f in con.execute(
        SQL_SENALES + " WHERE s.municipio_id = ? AND s.estado != 'Descartada' ORDER BY s.fecha DESC", (mid,))]
    boletin = leer_cache(con, f"boletin:{mid}")
    con_senal = sorted([v for v in variables if v["com_n"]], key=lambda v: -(v["com_puntaje"] or 0))
    nombres = nombres_escenarios(con)
    con.close()
    return render_template("municipio.html", m=m, tecnica=tecnica, comunidad=comunidad, senales=senales,
                           con_senal=con_senal, escenarios=nombres, ver_sensible=puede_ver_sensible(), boletin=boletin)


@app.route("/municipio/<mid>/boletin", methods=["POST"])
@requiere_login
def municipio_boletin(mid):
    con = conectar()
    m = con.execute("SELECT * FROM municipios WHERE id = ?", (mid,)).fetchone()
    if not m:
        con.close()
        abort(404)
    m, variables = resumen_municipio(con, m)
    datos = datos_boletin(con, variables, mid)
    contenido, motor = ia.boletin(con, f"{T()['Unidad']} {m['nombre']} ({m['departamento']}), {T()['tipo_unidad'].lower()}: {m['cultivo']}"
                                  + (f", {T()['marca_larga'].lower()}" if m["pdet"] else ""), datos)
    guardar_cache(con, f"boletin:{mid}", contenido, motor)
    registrar(con, "Boletín generado", f"{m['nombre']} · {'IA' if motor == 'ia' else 'modo básico'}")
    con.commit()
    con.close()
    flash("Boletín listo." if motor == "ia" else "Boletín generado en modo básico (sin IA).", "success")
    return redirect(url_for("municipio", mid=mid) + "#boletin")


@app.route("/municipio/<mid>/asamblea")
@requiere_login
def asamblea(mid):
    con = conectar()
    m = con.execute("SELECT * FROM municipios WHERE id = ?", (mid,)).fetchone()
    if not m:
        con.close()
        abort(404)
    m, variables = resumen_municipio(con, m)
    comunidad = analisis.indices_factores(variables, "com_puntaje")
    preocupaciones = [v for v in sorted(variables, key=lambda v: -(v["com_puntaje"] or 0))
                      if v["com_n"] and v["tipo"] == "R"][:3]
    oportunidades = [v for v in variables if v["tipo"] == "O"][:3]
    voces = con.execute("""SELECT s.descripcion FROM senales s LEFT JOIN variables v ON v.id = s.variable_id
                           WHERE s.municipio_id = ? AND s.estado = 'Validada' AND COALESCE(v.sensible, 0) = 0
                           AND s.descripcion != '' ORDER BY s.fecha DESC LIMIT 3""", (mid,)).fetchall()
    boletin = leer_cache(con, f"boletin:{mid}")
    compromisos = con.execute("""SELECT d.*, v.decision, v.nombre FROM decisiones d JOIN variables v ON v.id = d.variable_id
                                 WHERE d.estado = 'Aprobada' ORDER BY d.fecha_limite""").fetchall()
    con.close()
    return render_template("asamblea.html", m=m, comunidad=comunidad, preocupaciones=preocupaciones,
                           oportunidades=oportunidades, voces=voces, compromisos=compromisos, boletin=boletin)


# ---------------------------------------------------------------- voz del territorio
@app.route("/voz")
@requiere_login
def voz():
    con, variables, nombres = contexto()
    filas = [senal_dict(f) for f in con.execute(SQL_SENALES + " ORDER BY s.fecha DESC").fetchall()]
    mas, menos = analisis.brechas(variables)
    emergentes = grupos_emergentes(con, ver_sensible=puede_ver_sensible())
    con.close()
    return render_template("voz.html", por_validar=[f for f in filas if f["estado"] == "Por validar"],
                           validadas=[f for f in filas if f["estado"] == "Validada"],
                           descartadas=[f for f in filas if f["estado"] == "Descartada"],
                           brechas_mas=mas, brechas_menos=menos, escenarios=nombres,
                           ver_sensible=puede_ver_sensible(),
                           enlace=url_for("reportar", _external=True), umbral=analisis.UMBRAL_BRECHA,
                           emergentes=emergentes)


@app.route("/voz/<int:sid>/<accion>", methods=["POST"])
@requiere_login
@requiere_rol("analista")
def voz_validar(sid, accion):
    if accion not in ("validar", "descartar"):
        abort(404)
    estado = "Validada" if accion == "validar" else "Descartada"
    con = conectar()
    con.execute("UPDATE senales SET estado = ?, validado_por = ? WHERE id = ?", (estado, session["nombre"], sid))
    registrar(con, f"Señal {estado.lower()}", f"Señal #{sid}")
    con.commit()
    con.close()
    flash(f"Señal {estado.lower()}." + (" Ya cuenta en la lectura comunitaria." if estado == "Validada" else ""), "success")
    return redirect(url_for("voz"))


# ---------------------------------------------------------------- formulario comunitario (público)
@app.route("/reportar")
def reportar():
    con = conectar()
    lista = municipios(con)
    temas = temas_publicos(con)
    con.close()
    return render_template("reportar.html", municipios=lista, temas=temas)


def temas_publicos(con):
    return [dict(t) for t in con.execute(
        "SELECT id, senal, nombre, icono, sensible, tipo FROM variables WHERE senal IS NOT NULL ORDER BY tipo DESC, factor")]


def leer_foto(dato):
    """Recibe un data URL JPEG generado por el celular (ya sin metadatos GPS). Devuelve bytes o None."""
    if not dato or not isinstance(dato, str) or not dato.startswith("data:image/jpeg;base64,"):
        return None
    try:
        crudo = base64.b64decode(dato.split(",", 1)[1], validate=True)
    except Exception:
        return None
    return crudo if crudo[:3] == b"\xff\xd8\xff" and len(crudo) < 1_500_000 else None


def clave_sugerencia(texto, foto):
    h = hashlib.sha256((texto or "").strip().encode() + (hashlib.sha256(foto).digest() if foto else b""))
    return "sug:" + h.hexdigest()[:32]


@app.route("/api/sugerir", methods=["POST"])
def api_sugerir():
    """La IA lee el reporte y sugiere tema, gravedad y protege datos personales. No guarda la señal."""
    datos = request.get_json(silent=True) or {}
    usados = session.get("sugerencias", 0)
    if usados >= 40:
        return jsonify(ok=False, error="Demasiadas consultas seguidas. Elija el tema usted mismo."), 429
    session["sugerencias"] = usados + 1
    texto = (datos.get("texto") or "").strip()[:600]
    foto = leer_foto(datos.get("foto"))
    if len(texto) < 4 and not foto:
        return jsonify(ok=False, error="Cuéntenos un poco más."), 400
    con = conectar()
    temas = temas_publicos(con)
    r = ia.analizar_reporte(con, texto, temas, emergentes_existentes(con), foto)
    guardar_cache(con, clave_sugerencia(texto, foto), r, r["motor"])
    con.commit()
    con.close()
    tema = next((t for t in temas if t["id"] == r["tema"]), None)
    return jsonify(ok=True, tema=r["tema"], tema_nombre=tema["senal"] if tema else (r.get("tema_nuevo") or "Otro tema"),
                   icono=tema["icono"] if tema else "fa-lightbulb", gravedad=r["gravedad"], confianza=r["confianza"],
                   explicacion=r.get("explicacion"), texto_protegido=r["texto_protegido"],
                   protegidos=r["datos_protegidos"], tema_nuevo=r.get("tema_nuevo"), foto=r.get("foto"),
                   sensible=r["sensible"], motor=r["motor"])


@app.route("/api/senal", methods=["POST"])
def api_senal():
    datos = request.get_json(silent=True) or request.form
    if datos.get("sitio_web"):  # campo trampa contra robots
        return jsonify(ok=True)
    enviados = session.get("enviados", 0)
    if enviados >= 15:
        return jsonify(ok=False, error="Se alcanzó el límite de reportes por ahora. Intente más tarde."), 429
    try:
        gravedad = int(datos.get("gravedad", 0))
    except ValueError:
        gravedad = 0
    texto = (datos.get("descripcion") or datos.get("texto") or "").strip()[:600]
    foto = leer_foto(datos.get("foto"))
    con = conectar()
    mun = con.execute("SELECT id FROM municipios WHERE id = ?", (datos.get("municipio", ""),)).fetchone()
    tema_id = datos.get("tema", "")
    var = con.execute("SELECT id, sensible FROM variables WHERE id = ?", (tema_id,)).fetchone()
    if not mun or not (var or tema_id == "otro") or not 1 <= gravedad <= 5:
        con.close()
        return jsonify(ok=False, error="Faltan datos: elija municipio, tema y qué tanto afecta."), 400

    # Protección de datos en el servidor: nunca se confía en lo que llega del celular.
    analisis_ia = leer_cache(con, clave_sugerencia(texto, foto)) if texto or foto else None
    if (texto or foto) and not analisis_ia:  # p. ej. reportes guardados sin señal
        analisis_ia = ia.analizar_reporte(con, texto, temas_publicos(con), emergentes_existentes(con), foto)
    sensible = bool(var and var["sensible"])
    if analisis_ia:
        protegido, _ = ia.proteger_reglas(analisis_ia["texto_protegido"], sensible)
        analisis_ia["datos_protegidos"] = ia.contar_protegidos(protegido)
        analisis_ia["tema_persona"] = tema_id
    else:
        protegido = ""
    emergente = None
    if tema_id == "otro":
        emergente = (analisis_ia or {}).get("tema_nuevo") or "Otro tema (sin clasificar)"
    quien = datos.get("quien") if datos.get("quien") in T()["quien"] else "Prefiero no decir"
    if sensible:
        quien = "Prefiero no decir"  # los reportes de seguridad siempre son anónimos
    nombre_foto = None
    if foto:
        nombre_foto = f"{uuid.uuid4().hex}.jpg"
        with open(os.path.join(FOTOS_DIR, nombre_foto), "wb") as fh:
            fh.write(foto)
    meta = None
    if analisis_ia:
        meta = json.dumps({k: analisis_ia.get(k) for k in
                           ("motor", "tema", "tema_persona", "confianza", "explicacion", "datos_protegidos", "tema_nuevo", "foto")},
                          ensure_ascii=False)
    con.execute("""INSERT INTO senales (fecha, municipio_id, variable_id, gravedad, descripcion, quien, estado, origen,
                   emergente, foto, ia) VALUES (?, ?, ?, ?, ?, ?, 'Por validar', ?, ?, ?, ?)""",
                (datos.get("fecha_local") or ahora(), mun["id"], var["id"] if var else None, gravedad, protegido[:280],
                 quien, "Sin conexión" if datos.get("diferido") else "Formulario", emergente, nombre_foto, meta))
    con.execute("INSERT INTO bitacora (fecha, usuario, accion, detalle) VALUES (?, ?, ?, ?)",
                (ahora(), "reporte anónimo", "Señal recibida", f"{mun['id']} · {var['id'] if var else emergente}"))
    con.commit()
    con.close()
    session["enviados"] = enviados + 1
    return jsonify(ok=True, protegidos=len((analisis_ia or {}).get("datos_protegidos", [])))


@app.route("/foto/<nombre>")
@requiere_login
def foto(nombre):
    if not nombre.endswith(".jpg") or "/" in nombre or ".." in nombre:
        abort(404)
    return send_from_directory(FOTOS_DIR, nombre, mimetype="image/jpeg")


@app.route("/sw.js")
def service_worker():
    resp = send_from_directory(os.path.join(app.root_path, "static", "js"), "sw.js", mimetype="application/javascript")
    resp.headers["Cache-Control"] = "no-cache"
    return resp


@app.route("/manifest.webmanifest")
def manifest():
    return send_from_directory(os.path.join(app.root_path, "static"), "manifest.webmanifest",
                               mimetype="application/manifest+json")


# ---------------------------------------------------------------- variables
CAMPOS_ANALISTA = ["factor", "nombre", "descripcion", "tipo", "probabilidad", "impacto", "tendencia", "horizonte",
                   "cultivo", "fuente", "confiabilidad", "frecuencia", "decision", "responsable", "senal"]


def leer_form_variable(f):
    datos = {c: f.get(c) for c in CAMPOS_ANALISTA}
    datos["probabilidad"] = int(datos["probabilidad"])
    datos["impacto"] = int(datos["impacto"])
    return datos


@app.route("/variables")
@requiere_login
def variables():
    con, todas, nombres = contexto()
    filtro = request.args.get("factor", "todos")
    orden = {k: i for i, (k, _) in enumerate(FACTORES)}
    todas = sorted(todas, key=lambda v: (orden[v["factor"]], v["id"]))
    lista = todas if filtro == "todos" else [v for v in todas if v["factor"] == filtro]
    con.close()
    return render_template("variables.html", variables=lista, filtro=filtro, total=len(todas), escenarios=nombres)


@app.route("/variables/nueva", methods=["GET", "POST"])
@requiere_login
@requiere_rol("analista")
def variable_nueva():
    if request.method == "POST":
        con = conectar()
        d = leer_form_variable(request.form)
        n = con.execute("SELECT COUNT(*) FROM variables WHERE factor = ?", (d["factor"],)).fetchone()[0] + 1
        vid = f"{d['factor']}{n}"
        while con.execute("SELECT 1 FROM variables WHERE id = ?", (vid,)).fetchone():
            n += 1
            vid = f"{d['factor']}{n}"
        con.execute(f"""INSERT INTO variables (id, {', '.join(CAMPOS_ANALISTA)}, ods, icono, sensible, actualizado)
                        VALUES (?, {', '.join('?' * len(CAMPOS_ANALISTA))}, '', 'fa-circle-dot', 0, ?)""",
                    (vid, *[d[c] for c in CAMPOS_ANALISTA], ahora()))
        registrar(con, "Variable registrada", f"{vid} · {d['nombre']}" + nota_panel(request.form))
        origen = request.form.get("emergente")
        if origen:
            n = con.execute("UPDATE senales SET variable_id = ? WHERE emergente = ? AND variable_id IS NULL",
                            (vid, origen)).rowcount
            registrar(con, "Tema emergente convertido", f"«{origen}» → {vid} ({n} señales)")
        con.commit()
        con.close()
        flash(f"Variable «{d['nombre']}» registrada." + (" Las señales del tema emergente ya cuentan en ella." if origen else ""),
              "success")
        return redirect(url_for("variables"))
    propuesta = session.pop("propuesta_variable", None)
    return render_template("variable_form.html", v=propuesta, nueva=True)


@app.route("/variables/<vid>/editar", methods=["GET", "POST"])
@requiere_login
@requiere_rol("analista", "ti")
def variable_editar(vid):
    con = conectar()
    v = con.execute("SELECT * FROM variables WHERE id = ?", (vid,)).fetchone()
    if not v:
        con.close()
        flash("La variable no existe.", "warning")
        return redirect(url_for("variables"))
    if request.method == "POST":
        f = request.form
        if session["rol"] == "analista":
            d = leer_form_variable(f)
            con.execute(f"""UPDATE variables SET {', '.join(c + '=?' for c in CAMPOS_ANALISTA)}, actualizado=?
                            WHERE id=?""", (*[d[c] for c in CAMPOS_ANALISTA], ahora(), vid))
            detalle = f"{vid} · P{v['probabilidad']}×I{v['impacto']} → P{d['probabilidad']}×I{d['impacto']}" + nota_panel(f)
        else:
            con.execute("UPDATE variables SET fuente=?, confiabilidad=?, frecuencia=?, actualizado=? WHERE id=?",
                        (f.get("fuente"), f["confiabilidad"], f["frecuencia"], ahora(), vid))
            detalle = f"{vid} · fuente {f['confiabilidad']}, {f['frecuencia']}"
        registrar(con, "Variable actualizada", detalle)
        con.commit()
        con.close()
        flash("Cambios guardados. El nivel se recalculó.", "success")
        return redirect(url_for("variables"))
    con.close()
    return render_template("variable_form.html", v=dict(v))


def nota_panel(form):
    """Si la calificación salió del panel de expertos, se deja constancia en la bitácora."""
    acuerdo = form.get("panel_acuerdo", "")
    return f" · calificación del panel de expertos (acuerdo {acuerdo.lower()})" if acuerdo in ("Alto", "Medio", "Bajo") else ""


@app.route("/ia/panel", methods=["POST"])
@requiere_login
@requiere_rol("analista")
def ia_panel():
    """Panel de expertos: seis perspectivas PESTEL califican la variable por separado (método Delphi)."""
    f = request.form
    variable = {"factor": f.get("factor"), "nombre": f.get("nombre", "").strip(), "descripcion": f.get("descripcion", "").strip(),
                "tipo": f.get("tipo"), "horizonte": f.get("horizonte")}
    if len(variable["nombre"]) < 4:
        return jsonify(ok=False, error="Escriba primero el nombre de la variable y, si puede, por qué afecta a la organización.")
    con = conectar()
    opiniones, error = ia.panel_expertos(con, variable)
    if opiniones:
        c = analisis.consenso(opiniones)
        registrar(con, "Panel de expertos consultado", f"{variable['nombre'][:60]} · P{c['probabilidad']}×I{c['impacto']}, acuerdo {c['acuerdo'].lower()}")
        con.commit()
    con.close()
    if error:
        return jsonify(ok=False, error=error)
    return jsonify(ok=True, opiniones=opiniones, consenso=analisis.consenso(opiniones))


@app.route("/variables/<vid>/eliminar", methods=["POST"])
@requiere_login
@requiere_rol("analista")
def variable_eliminar(vid):
    con = conectar()
    v = con.execute("SELECT nombre FROM variables WHERE id = ?", (vid,)).fetchone()
    if v:
        con.execute("DELETE FROM variables WHERE id = ?", (vid,))
        registrar(con, "Variable eliminada", f"{vid} · {v['nombre']}")
        con.commit()
        flash(f"Variable «{v['nombre']}» eliminada.", "info")
    con.close()
    return redirect(url_for("variables"))


@app.route("/emergentes/variable", methods=["POST"])
@requiere_login
@requiere_rol("analista")
def emergente_a_variable():
    """La IA propone cómo registrar un tema emergente como variable PESTEL; la técnica revisa y guarda."""
    tema = request.form.get("tema", "")
    con = conectar()
    textos = [r[0] for r in con.execute("SELECT descripcion FROM senales WHERE emergente = ? AND variable_id IS NULL "
                                        "AND descripcion != ''", (tema,))]
    if not textos:
        con.close()
        flash("Ese tema ya no tiene señales pendientes.", "warning")
        return redirect(url_for("voz"))
    propuesta, motor = ia.proponer_variable(con, tema, textos)
    con.close()
    propuesta["emergente"] = tema
    propuesta["motor"] = motor
    propuesta.update(fuente=f"Reportes de {T()['gente']} (Prisma)", confiabilidad="Media", frecuencia="Semanal",
                     tendencia="Sube", horizonte="Corto")
    session["propuesta_variable"] = propuesta
    return redirect(url_for("variable_nueva"))


@app.route("/ajustes/ia", methods=["GET", "POST"])
@requiere_login
@requiere_rol("ti")
def ajustes_ia():
    con = conectar()
    prueba = None
    if request.method == "POST":
        accion = request.form.get("accion")
        if accion == "guardar":
            clave = request.form.get("clave", "").strip()
            if not clave.startswith("sk-"):
                flash("La clave debe empezar por «sk-». Cópiela completa desde console.anthropic.com.", "warning")
            else:
                con.execute("INSERT OR REPLACE INTO ajustes (clave, valor) VALUES ('clave_ia', ?)", (clave,))
                registrar(con, "Clave de IA configurada", "termina en …" + clave[-4:])
                con.commit()
                prueba = ia.probar(con)
        elif accion == "borrar":
            con.execute("DELETE FROM ajustes WHERE clave = 'clave_ia'")
            registrar(con, "Clave de IA eliminada")
            con.commit()
            flash("Clave eliminada. Prisma sigue funcionando en modo básico.", "info")
        elif accion == "probar":
            prueba = ia.probar(con)
    clave = ia.obtener_clave(con)
    desde_entorno = bool(os.environ.get("ANTHROPIC_API_KEY"))
    uso = ia.uso_hoy(con)
    con.close()
    return render_template("ajustes_ia.html", configurada=bool(clave), final=clave[-4:] if clave else "",
                           desde_entorno=desde_entorno, prueba=prueba, modelo=ia.MODELO, error=ia.ultimo_error["texto"],
                           uso=uso, tope=ia.TOPE_DIARIO)


@app.route("/seguimiento")
def seguimiento():
    """Página pública: qué se ha hecho con lo que reportó la comunidad. Sin datos personales."""
    con = conectar()
    compromisos = con.execute("""SELECT d.*, v.decision, v.nombre, v.senal, v.icono FROM decisiones d
                                 JOIN variables v ON v.id = d.variable_id WHERE d.estado = 'Aprobada'
                                 ORDER BY d.fecha_limite""").fetchall()
    temas = con.execute("""SELECT COALESCE(v.senal, s.emergente) AS tema, COALESCE(v.icono, 'fa-lightbulb') AS icono,
                           COUNT(*) AS n FROM senales s LEFT JOIN variables v ON v.id = s.variable_id
                           WHERE s.estado = 'Validada' GROUP BY 1, 2 ORDER BY n DESC LIMIT 8""").fetchall()
    total = con.execute("SELECT COUNT(*) FROM senales WHERE estado = 'Validada'").fetchone()[0]
    con.close()
    return render_template("seguimiento.html", compromisos=compromisos, temas=temas, total=total)


# ---------------------------------------------------------------- escenarios
@app.route("/escenarios")
@requiere_login
def escenarios():
    con, variables, nombres = contexto()
    lista = []
    for e in con.execute("SELECT * FROM escenarios").fetchall():
        afectadas = [r["nombre"] for r in con.execute(
            """SELECT v.nombre FROM escenario_efectos ef JOIN variables v ON v.id = ef.variable_id
               WHERE ef.escenario_id = ?""", (e["id"],)).fetchall()]
        lista.append({**dict(e), "afectadas": afectadas, "activo": e["id"] in escenarios_activos(),
                      "cadena": ia.limpiar_cadena(json.loads(e["cadena"] or "[]"))})
    entrando = {x["id"]: x for x in analisis.escenarios_en_entrada(con)}
    factores = analisis.indices_factores(variables, "puntaje")
    base = analisis.indices_factores(variables, "puntaje_base")
    cambiadas = sorted([v for v in variables if v["modificada"]], key=lambda v: -v["puntaje"])
    con.close()
    return render_template("escenarios.html", lista=lista, factores=factores, base=base, cambiadas=cambiadas,
                           entrando=entrando, entrada_dias=analisis.ENTRADA_DIAS, entrada_min=analisis.ENTRADA_MINIMO,
                           escenarios=nombres)


@app.route("/escenarios/proponer", methods=["POST"])
@requiere_login
@requiere_rol("analista", "directivo")
def escenario_proponer():
    """La IA propone un escenario a partir de una pregunta «¿Y si…?». No se guarda hasta que una persona lo revise."""
    pregunta = request.form.get("pregunta", "").strip()
    if len(pregunta) < 6:
        flash("Escriba la pregunta completa, por ejemplo: «¿y si hay un paro camionero de dos semanas?».", "warning")
        return redirect(url_for("escenarios"))
    con = conectar()
    variables = analisis.evaluar(con, ())
    propuesta, motor = ia.proponer_escenario(con, pregunta, variables)
    registrar(con, "Escenario propuesto", f"«{pregunta[:80]}» · {'IA' if motor == 'ia' else 'modo básico'}")
    con.commit()
    con.close()
    return render_template("escenario_nuevo.html", propuesta=propuesta, motor=motor, variables=variables)


@app.route("/escenarios/crear", methods=["POST"])
@requiere_login
@requiere_rol("analista", "directivo")
def escenario_crear():
    f = request.form
    nombre = f.get("nombre", "").strip()[:80]
    if not nombre:
        flash("El escenario necesita un nombre.", "warning")
        return redirect(url_for("escenarios"))
    con = conectar()
    ids = {r[0] for r in con.execute("SELECT id FROM variables")}
    efectos = []
    for vid in f.getlist("variable"):
        if vid in ids and f.get(f"incluir_{vid}"):
            try:
                p, i = int(f.get(f"p_{vid}")), int(f.get(f"i_{vid}"))
            except (TypeError, ValueError):
                continue
            efectos.append((vid, min(5, max(1, p)), min(5, max(1, i))))
    if not efectos:
        con.close()
        flash("Marque al menos una variable para el escenario.", "warning")
        return redirect(url_for("escenarios"))
    eid = "x" + uuid.uuid4().hex[:8]
    try:
        cadena = ia.limpiar_cadena(json.loads(f.get("cadena") or "[]"))
    except ValueError:
        cadena = []
    con.execute("""INSERT INTO escenarios (id, nombre, descripcion, origen, pregunta, creado_por, cadena)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (eid, nombre, f.get("descripcion", "").strip()[:300], "reglas" if f.get("motor") == "reglas" else "ia",
                 f.get("pregunta", "")[:300], session["nombre"], json.dumps(cadena, ensure_ascii=False)))
    con.executemany("INSERT INTO escenario_efectos VALUES (?, ?, ?, ?)", [(eid, v, p, i) for v, p, i in efectos])
    registrar(con, "Escenario creado", f"{nombre} · {len(efectos)} variables")
    con.commit()
    con.close()
    session["escenarios"] = escenarios_activos() + [eid]
    flash(f"Escenario «{nombre}» guardado y activado. Mire el impacto en el entorno.", "success")
    return redirect(url_for("escenarios") + "#t-impacto")


@app.route("/escenarios/<eid>/eliminar", methods=["POST"])
@requiere_login
@requiere_rol("analista", "directivo")
def escenario_eliminar(eid):
    con = conectar()
    e = con.execute("SELECT nombre, origen FROM escenarios WHERE id = ?", (eid,)).fetchone()
    if e and e["origen"] != "base":
        con.execute("DELETE FROM escenario_efectos WHERE escenario_id = ?", (eid,))
        con.execute("DELETE FROM escenarios WHERE id = ?", (eid,))
        registrar(con, "Escenario eliminado", e["nombre"])
        con.commit()
        session["escenarios"] = [x for x in escenarios_activos() if x != eid]
        flash(f"Escenario «{e['nombre']}» eliminado.", "info")
    con.close()
    return redirect(url_for("escenarios"))


@app.route("/escenarios/<eid>/alternar", methods=["POST"])
@requiere_login
def escenario_alternar(eid):
    activos = escenarios_activos()
    session["escenarios"] = [e for e in activos if e != eid] if eid in activos else activos + [eid]
    if request.form.get("ir") == "impacto":
        return redirect(url_for("escenarios") + "#t-impacto")
    return redirect(request.referrer or url_for("escenarios"))


@app.route("/escenarios/limpiar", methods=["POST"])
@requiere_login
def escenarios_limpiar():
    session["escenarios"] = []
    return redirect(request.referrer or url_for("inicio"))


# ---------------------------------------------------------------- decisiones y compromisos
@app.route("/decisiones")
@requiere_login
def decisiones():
    con, variables, nombres = contexto()
    estados = {r["variable_id"]: dict(r) for r in con.execute("SELECT * FROM decisiones").fetchall()}
    items = analisis.decisiones_sugeridas(variables)
    abierta = lambda v: estados.get(v["id"], {}).get("estado", "Pendiente") in ("Pendiente", "En análisis")
    por_id = {v["id"]: v for v in variables}
    compromisos = [por_id[k] for k, e in estados.items() if e["estado"] == "Aprobada" and k in por_id]
    descartadas = [por_id[k] for k, e in estados.items() if e["estado"] == "Descartada" and k in por_id]
    con.close()
    return render_template("decisiones.html", estados=estados, escenarios=nombres,
                           por_revisar=[v for v in items if abierta(v)], compromisos=compromisos,
                           descartadas=descartadas)


@app.route("/decisiones/<vid>", methods=["POST"])
@requiere_login
@requiere_rol("directivo")
def decision_actualizar(vid):
    con = conectar()
    f = request.form
    estado = f.get("estado", "Pendiente")
    actores = ", ".join(a for a in f.getlist("actores") if a in g.P["actores"])
    con.execute(
        """INSERT INTO decisiones (variable_id, estado, comentario, actores, fecha_limite, avance, actualizado_por, fecha)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)
           ON CONFLICT(variable_id) DO UPDATE SET estado=excluded.estado, comentario=excluded.comentario,
           actores=excluded.actores, fecha_limite=excluded.fecha_limite, avance=COALESCE(decisiones.avance, excluded.avance),
           actualizado_por=excluded.actualizado_por, fecha=excluded.fecha""",
        (vid, estado, f.get("comentario", "").strip(), actores, f.get("fecha_limite") or None, "Sin iniciar",
         session["nombre"], ahora()))
    registrar(con, "Decisión actualizada", f"{vid} → {estado}" + (f" · {actores}" if actores else ""))
    con.commit()
    con.close()
    flash("Compromiso registrado." if estado == "Aprobada" else f"Decisión marcada como «{estado}».", "success")
    return redirect(url_for("decisiones"))


@app.route("/decisiones/<vid>/avance", methods=["POST"])
@requiere_login
@requiere_rol("directivo", "analista")
def decision_avance(vid):
    avance = request.form.get("avance")
    if avance not in OPCIONES["avance"]:
        abort(400)
    con = conectar()
    con.execute("UPDATE decisiones SET avance = ? WHERE variable_id = ?", (avance, vid))
    registrar(con, "Avance de compromiso", f"{vid} → {avance}")
    con.commit()
    con.close()
    flash(f"Avance actualizado: {avance}.", "success")
    return redirect(url_for("decisiones"))


@app.route("/decisiones/<vid>/reabrir", methods=["POST"])
@requiere_login
@requiere_rol("directivo")
def decision_reabrir(vid):
    con = conectar()
    con.execute("DELETE FROM decisiones WHERE variable_id = ?", (vid,))
    registrar(con, "Decisión reabierta", f"{vid} → Pendiente")
    con.commit()
    con.close()
    flash("La decisión volvió a «Por revisar».", "info")
    return redirect(url_for("decisiones"))


# ---------------------------------------------------------------- informe
@app.route("/informe")
@requiere_login
def informe():
    con, variables, nombres = contexto()
    estados = {r["variable_id"]: dict(r) for r in con.execute("SELECT * FROM decisiones").fetchall()}
    factores = analisis.indices_factores(variables, "efectivo")
    tecnica = analisis.indices_factores(variables, "puntaje")
    comunidad = analisis.indices_factores(variables, "com_puntaje")
    riesgos = sorted([v for v in variables if v["tipo"] == "R"], key=lambda v: -v["efectivo"])[:5]
    oportunidades = sorted([v for v in variables if v["tipo"] == "O"], key=lambda v: -v["efectivo"])
    sugeridas = analisis.decisiones_sugeridas(variables)
    mas, menos = analisis.brechas(variables)
    total_senales = con.execute("SELECT COUNT(*) FROM senales WHERE estado = 'Validada'").fetchone()[0]
    lista_m = [resumen_municipio(con, m)[0] for m in municipios(con)]
    boletin = leer_cache(con, "boletin:general")
    emergentes = grupos_emergentes(con)
    con.close()
    return render_template("informe.html", factores=factores, tecnica=tecnica, comunidad=comunidad, riesgos=riesgos,
                           oportunidades=oportunidades, sugeridas=sugeridas, estados=estados, escenarios=nombres,
                           brechas_mas=mas, brechas_menos=menos,
                           total_senales=total_senales, municipios=lista_m, boletin=boletin, emergentes=emergentes)


@app.route("/informe/boletin", methods=["POST"])
@requiere_login
def informe_boletin():
    con, variables, _ = contexto()
    contenido, motor = ia.boletin(con, f"Toda la organización: {T()['sector']}", datos_boletin(con, variables))
    guardar_cache(con, "boletin:general", contenido, motor)
    registrar(con, "Resumen ejecutivo generado", "IA" if motor == "ia" else "modo básico")
    con.commit()
    con.close()
    flash("Resumen ejecutivo listo." if motor == "ia" else "Resumen generado en modo básico (sin IA).", "success")
    return redirect(url_for("informe"))


@app.route("/informe/csv")
@requiere_login
def informe_csv():
    con, variables, _ = contexto()
    con.close()
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["ID", "Factor", "Variable", "Tipo", T()["categoria"], "Lectura técnica", T()["lectura_gente"],
                "Señales", "Brecha", "Nivel", "Fuente", "Decisión sugerida", "Responsable"])
    for v in variables:
        w.writerow([v["id"], v["factor_nombre"], v["nombre"], "Riesgo" if v["tipo"] == "R" else "Oportunidad",
                    v["cultivo"], v["puntaje"], v["com_puntaje"] or "", v["com_n"], v["brecha"] or "",
                    v["nivel"], v["fuente"], v["decision"], v["responsable"]])
    nombre = f"{APP_NOMBRE.lower()}_{g.P['id']}_{hora_local():%Y%m%d}.csv"
    return Response("﻿" + buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


@app.route("/informe/excel")
@requiere_login
def informe_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    con, variables, nombres = contexto()
    senales = con.execute("""SELECT s.fecha, m.nombre AS municipio, m.departamento,
                             COALESCE(v.nombre, s.emergente || ' (tema emergente)') AS variable,
                             COALESCE(v.sensible, 0) AS sensible, s.gravedad, s.descripcion, s.quien, s.estado FROM senales s
                             LEFT JOIN variables v ON v.id = s.variable_id JOIN municipios m ON m.id = s.municipio_id
                             ORDER BY s.fecha DESC""").fetchall()
    con.close()
    oscuro = PatternFill("solid", fgColor="143D34")
    colores = {"Crítico": "F7DEDA", "Alto": "FBE9DA", "Medio": "F8F0D6", "Bajo": "E4F2E8"}
    borde = Border(bottom=Side(style="thin", color="E2E6DF"))
    blanco = Font(bold=True, color="FFFFFF")

    def encabezado(ws, fila):
        for celda in ws[fila]:
            celda.fill, celda.font = oscuro, blanco
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Lectura PESTEL"
    ws["A1"] = f"{APP_NOMBRE} · {APP_LEMA} · {T()['sector']}"
    ws["A1"].font = Font(bold=True, size=14, color="143D34")
    ws["A2"] = (f"Generado el {hora_local():%d/%m/%Y %H:%M} por {session['nombre']}"
                + (f" · Escenario: {' + '.join(nombres)}" if nombres else " · Entorno actual") + " · Datos de demostración")
    ws["A2"].font = Font(italic=True, color="5E6B63")
    ws.append([])
    ws.append(["ID", "Factor", "Variable", "Tipo", T()["categoria"], "Lectura técnica", T()["lectura_gente"],
               "Señales", "Brecha", "Nivel", "Fuente", "Decisión sugerida", "Responsable"])
    encabezado(ws, 4)
    orden = {k: i for i, (k, _) in enumerate(FACTORES)}
    for v in sorted(variables, key=lambda v: (orden[v["factor"]], v["id"])):
        es_r = v["tipo"] == "R"
        ws.append([v["id"], v["factor_nombre"], v["nombre"], "Riesgo" if es_r else "Oportunidad", v["cultivo"],
                   v["puntaje"], v["com_puntaje"], v["com_n"], v["brecha"],
                   v["nivel"] if es_r else "Oportunidad", v["fuente"], v["decision"], v["responsable"]])
        fila = ws.max_row
        color = colores.get(v["nivel"], "FFFFFF") if es_r else "E3EEF6"
        ws[f"K{fila}"].fill = PatternFill("solid", fgColor=color)
        ws[f"K{fila}"].font = Font(bold=True)
        if es_r and v["brecha"] is not None and abs(v["brecha"]) >= analisis.UMBRAL_BRECHA:
            ws[f"J{fila}"].font = Font(bold=True, color="A3281E")
        for celda in ws[fila]:
            celda.border = borde
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    for i, ancho in enumerate([6, 12, 34, 12, 9, 8, 10, 12, 9, 8, 11, 28, 44, 22]):
        ws.column_dimensions[chr(65 + i)].width = ancho
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:N{ws.max_row}"

    hs = wb.create_sheet("Voz del territorio")
    hs.append(["Fecha", "Municipio", "Departamento", "Tema", "Qué tanto afecta (1-5)", "Lo que dijo", "Quién", "Estado"])
    encabezado(hs, 1)
    for s in senales:
        texto = "Reservado por seguridad" if s["sensible"] else (s["descripcion"] or "")
        hs.append([s["fecha"], s["municipio"], s["departamento"], s["variable"], s["gravedad"], texto,
                   "Anónimo" if s["sensible"] else s["quien"], s["estado"]])
    for col, ancho in zip("ABCDEFGH", [16, 22, 18, 34, 12, 48, 18, 12]):
        hs.column_dimensions[col].width = ancho

    rs = wb.create_sheet("Resumen por factor")
    rs.append(["Factor", "Lectura técnica", "Lectura comunitaria", "Nivel efectivo", "Señales"])
    encabezado(rs, 1)
    for f, t, c in zip(analisis.indices_factores(variables, "efectivo"), analisis.indices_factores(variables, "puntaje"),
                       analisis.indices_factores(variables, "com_puntaje")):
        rs.append([f["nombre"], t["indice"], c["indice"], f["nivel"], f["senales"]])
        rs.cell(rs.max_row, 4).fill = PatternFill("solid", fgColor=colores.get(f["nivel"], "FFFFFF"))
    for col, ancho in zip("ABCDE", [16, 16, 20, 14, 10]):
        rs.column_dimensions[col].width = ancho

    salida = io.BytesIO()
    wb.save(salida)
    nombre = f"{APP_NOMBRE.lower()}_{g.P['id']}_{hora_local():%Y%m%d}.xlsx"
    return Response(salida.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


# ---------------------------------------------------------------- otras vistas
@app.route("/contexto")
@requiere_login
def contexto_territorio():
    con = conectar()
    lista = municipios(con)
    con.close()
    return render_template("contexto.html", municipios=lista, contexto=g.P["contexto"], fuentes=g.P["fuentes"])


@app.route("/requerimientos")
@requiere_login
def requerimientos():
    return render_template("requerimientos.html")


@app.route("/bitacora")
@requiere_login
@requiere_rol("ti", "directivo")
def bitacora():
    con = conectar()
    filas = con.execute("SELECT * FROM bitacora ORDER BY id DESC LIMIT 300").fetchall()
    usuarios = con.execute("SELECT usuario, nombre, rol FROM usuarios").fetchall()
    con.close()
    return render_template("bitacora.html", filas=filas, usuarios=usuarios)


@app.route("/reiniciar", methods=["POST"])
@requiere_login
@requiere_rol("ti")
def reiniciar():
    inicializar(forzar=True)
    session["escenarios"] = []
    flash("Datos de demostración restablecidos.", "info")
    return redirect(url_for("bitacora"))


if __name__ == "__main__":
    inicializar()
    print("\n  Prisma listo en http://127.0.0.1:5000")
    print("  Formulario para reportar: http://127.0.0.1:5000/reportar\n")
    app.run(debug=os.environ.get("PRISMA_DEBUG") == "1")  # nunca en modo depuración en producción
